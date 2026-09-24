#!/usr/bin/env python3
"""轻量版排版：guide/<lang>.md → 16:9 幻灯片 HTML + PDF + PPTX + 网页用 JSON。
用法：.venv/bin/python build_guide.py zh [--url https://canlah.ai/playbook/]
guide md 格式：每页一个 `## 主张`；下面一个图围栏（DIAGRAM_SPEC 五种之一）、≤120 字正文、
`<!-- full: 完整版小节 -->`。第一页是封面，含 [QR] 的那页是「看完整版」页。
"""
import os, re, sys, json, base64, io, importlib.util, asyncio

HERE = os.path.dirname(os.path.abspath(__file__))
V3 = os.path.dirname(HERE)
OUT = os.path.join(V3, '_build')
spec = importlib.util.spec_from_file_location('b3', os.path.join(HERE, 'build_v3.py'))
b3 = importlib.util.module_from_spec(spec); spec.loader.exec_module(b3)

BOOK_CSS = re.search(r"CSS = r'''(.*?)'''", open(os.path.join(HERE, 'emit_v3.py')).read(), re.S).group(1)

LANG = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith('--') else 'zh'
URL = sys.argv[sys.argv.index('--url') + 1] if '--url' in sys.argv else (
    'https://canlah.ai/zh/playbook/' if LANG.startswith('zh') else 'https://canlah.ai/playbook/')
REPO = 'https://github.com/Canlah-AI/geo-playbook'
T = {
    'zh': {'kick': 'CANLAH AI · GEO PLAYBOOK 轻量版', 'scan': '扫码看完整版', 'agent': '给 AI 助手用的版本（开源）', 'page': '第 {} 页', 'full': '完整版 {} 节', 'fullapp': '完整版附录 {}'},
    'en': {'kick': 'CANLAH AI · GEO PLAYBOOK · QUICK GUIDE', 'scan': 'Scan for the full playbook', 'agent': 'For AI assistants (open source)', 'page': 'Page {}', 'full': 'Full playbook §{}', 'fullapp': 'Full playbook, appendix {}'},
}[LANG if LANG in ('zh', 'en') else 'zh']

def parse(md):
    md = re.sub(r'^#\s+.*\n', '', md, count=1)
    parts = re.split(r'^##\s+', md, flags=re.M)
    pages = []
    for p in parts[1:]:
        title, _, rest = p.partition('\n')
        full = re.search(r'<!--\s*full:\s*(.*?)\s*-->', rest)
        rest = re.sub(r'<!--.*?-->', '', rest, flags=re.S)
        fence = re.search(r'```.*?\n```', rest, flags=re.S)
        fig_md = fence.group(0) if fence else ''
        body_md = (rest[:fence.start()] + rest[fence.end():]) if fence else rest
        pages.append({'title': title.strip(), 'fig_md': fig_md, 'body_md': body_md.strip(),
                      'full': full.group(1) if full else '', 'qr': '[QR]' in rest})
    return pages

def full_label(s):
    # 页脚只印小节号，不印小节全名（全名里是完整版的内部术语）
    m = re.match(r'^\s*([A-Z])\.(\d+)', s or '')
    if m:
        return T['fullapp'].format(m.group(1) + '.' + m.group(2))
    m = re.match(r'^\s*(\d+(?:\.[\dA-Z]+)?)', s or '')
    return T['full'].format(m.group(1)) if m else ''

def render(md):
    if not md.strip():
        return ''
    html, _ = b3.convert(md.replace('[QR]', ''), 'guide', {})
    return html

def qr_png(url):
    import qrcode
    q = qrcode.QRCode(border=1, box_size=12, error_correction=qrcode.constants.ERROR_CORRECT_M)
    q.add_data(url); q.make(fit=True)
    img = q.make_image(fill_color='#1a1919', back_color='#ffffff')
    buf = io.BytesIO(); img.save(buf, format='PNG'); return buf.getvalue()

SLIDE_CSS = r'''
:root{--paper:#fafaf9;--sheet:#ffffff;--sheet-2:#f1f0ed;--ink:#1a1919;--ink-2:#57534e;--muted:#6f6b66;
  --rule:#e8e6e1;--rule-2:#d6d3cd;--accent:#4f46e5;--accent-soft:#eef0fd;--accent-ink:#3730a3;
  --warn:#b4531a;--warn-soft:#fbf0e6;
  --serif:"DM Sans","Noto Sans SC","PingFang SC",sans-serif;--sans:"DM Sans","Noto Sans SC","PingFang SC",sans-serif;
  --mono:"DM Mono",ui-monospace,Menlo,"Noto Sans SC","PingFang SC",monospace;}
@page{size:1920px 1080px;margin:0}
html,body{margin:0;background:#d9d6d0}
.slide{width:1920px;height:1080px;box-sizing:border-box;background:var(--paper);color:var(--ink);font-family:var(--sans);
  position:relative;overflow:hidden;page-break-after:always;break-after:page;margin:0 auto 24px;display:grid;
  grid-template-columns:720px 1fr;gap:72px;padding:120px 110px 110px}
@media print{html,body{background:var(--paper)}.slide{margin:0}}
.kick{position:absolute;top:52px;left:110px;right:110px;display:flex;justify-content:space-between;
  font-family:var(--mono);font-size:20px;letter-spacing:.14em;color:var(--muted)}
.kick b{color:var(--accent);font-weight:500}
.lhs{display:flex;flex-direction:column;justify-content:center;gap:34px;min-width:0}
.lhs h2{font-size:58px;line-height:1.22;font-weight:700;letter-spacing:-.01em;margin:0;text-wrap:balance}
.lhs .body{font-size:28px;line-height:1.72;color:var(--ink-2)}
.lhs .body p{margin:0 0 18px}
.lhs .body strong{color:var(--ink)}
.lhs .full{font-family:var(--mono);font-size:19px;color:var(--muted)}
.rhs{display:flex;align-items:center;justify-content:center;min-width:0}
.rhs .fig{margin:0;width:100%}
.rhs .fig.wf{--z:1.5}.rhs .fig.split{--z:1.55}.rhs .fig.bars{--z:1.7}.rhs .fig.steps{--z:1.45}
.rhs .fig.wf,.rhs .fig.split,.rhs .fig.bars,.rhs .fig.steps{zoom:var(--z);width:calc(100% / var(--z))}
.rhs .fig .sp,.rhs .fig .bar-list,.rhs .fig .wf-page{max-width:none}
.rhs .fig figcaption{display:none}
.rhs .mmd-out{border:none;background:none;padding:0}
.rhs .mmd-out{display:flex;justify-content:center;align-items:center;width:1000px;height:800px}
.rhs .mmd-out svg{display:block}
.cover{grid-template-columns:1fr;align-content:center;padding:0 180px;background:var(--ink);color:#fafaf9}
.cover .kick{color:#a8a29e}.cover .kick b{color:#b5bffa}
.cover h1{font-size:92px;line-height:1.14;margin:0 0 36px;font-weight:700;letter-spacing:-.02em;max-width:1400px;text-wrap:balance}
.cover .sub{font-size:34px;line-height:1.6;color:#d6d3cd;max-width:1300px}
.cover .cbrand{margin-top:70px;font-family:var(--mono);font-size:24px;color:#b5bffa;letter-spacing:.08em}
.qrs{grid-template-columns:1fr 520px;align-items:center}
.qrbox{background:#fff;border:1px solid var(--rule);border-radius:16px;padding:34px;text-align:center}
.qrbox img{width:420px;height:420px;display:block;margin:0 auto 18px}
.qrbox .u{font-family:var(--mono);font-size:22px;color:var(--accent-ink);word-break:break-all}
.qrbox .l{font-size:22px;color:var(--muted);margin-top:6px}
.repo{margin-top:26px;font-family:var(--mono);font-size:22px;color:var(--ink-2)}
.pn{position:absolute;bottom:46px;right:110px;font-family:var(--mono);font-size:18px;color:var(--faint,#a8a29e)}
'''

MMD_JS = r'''
(function(){
  function tok(n){return getComputedStyle(document.documentElement).getPropertyValue(n).trim();}
  function esc(s){return s.replace(/[&<>]/g,function(m){return {'&':'&amp;','<':'&lt;','>':'&gt;'}[m];});}
  var figs=[].slice.call(document.querySelectorAll('figure.mmd'));
  if(!window.mermaid||!figs.length){window.__ready=true;return;}
  mermaid.initialize({startOnLoad:false,securityLevel:'strict',theme:'base',fontFamily:tok('--sans'),
    themeVariables:{primaryColor:'#ffffff',primaryBorderColor:tok('--rule-2'),primaryTextColor:tok('--ink'),
      lineColor:tok('--muted'),textColor:tok('--ink-2'),edgeLabelBackground:tok('--paper'),fontSize:'24px'},
    flowchart:{htmlLabels:true,curve:'basis',padding:14,nodeSpacing:44,rankSpacing:52,wrappingWidth:200}});
  var defs='\n  classDef hl fill:'+tok('--accent-soft')+',stroke:'+tok('--accent')+',color:'+tok('--accent-ink')+',stroke-width:2px;'
          +'\n  classDef warn fill:'+tok('--warn-soft')+',stroke:'+tok('--warn')+',color:'+tok('--warn')+';';
  var i=0;
  figs.reduce(function(p,f){return p.then(function(){
    var src=f.querySelector('.mmd-src').textContent.replace(/(\b[A-Za-z_]\w*)\{"([^"{}]*)"\}/g,'$1{{"$2"}}');
    var out=f.querySelector('.mmd-out');
    function fit(){var s=out.querySelector('svg');if(!s)return 1;var vb=s.viewBox.baseVal,w=vb.width||1,h=vb.height||1;
      var sc=Math.min(1000/w,800/h,2.4);s.style.maxWidth='none';s.style.width=(w*sc)+'px';s.style.height=(h*sc)+'px';return w/h;}
    return mermaid.render('gm'+(++i),src+defs).then(function(r){out.innerHTML=r.svg;f.dataset.ok='1';
      var ar=fit();
      if(ar>2.2&&/^\s*flowchart\s+LR/.test(src)){return mermaid.render('gm'+(++i),src.replace(/^(\s*flowchart\s+)LR/,'$1TD')+defs).then(function(r2){out.innerHTML=r2.svg;fit();});}
    })
      .catch(function(e){f.querySelector('.mmd-out').innerHTML='<pre>'+esc(src)+'</pre>';f.dataset.err=String(e&&e.message||e);});
  });},Promise.resolve()).then(function(){window.__ready=true;});
})();
'''

def slide_html(pages, qr_b64):
    n = len(pages); out = []
    for i, p in enumerate(pages, 1):
        kick = '<div class="kick"><span><b>%s</b></span><span>%d / %d</span></div>' % (b3.esc(T['kick']), i, n)
        if i == 1:
            sub = render(p['body_md'])
            out.append('<section class="slide cover" data-i="%d">%s<div><h1>%s</h1><div class="sub">%s</div>'
                       '<div class="cbrand">canlah.ai</div></div></section>' % (i, kick, b3.inline(p['title']), sub))
            continue
        if p['qr']:
            out.append('<section class="slide qrs" data-i="%d">%s<div class="lhs"><h2>%s</h2><div class="body">%s</div>'
                       '<div class="repo">%s<br>%s</div></div>'
                       '<div class="rhs"><div class="qrbox"><img alt="QR" src="data:image/png;base64,%s"><div class="u">%s</div>'
                       '<div class="l">%s</div></div></div></section>' % (
                           i, kick, b3.inline(p['title']), render(p['body_md']), b3.esc(T['agent']), b3.esc(REPO),
                           qr_b64, b3.esc(URL.replace('https://', '')), b3.esc(T['scan'])))
            continue
        full = '<div class="full">→ %s</div>' % b3.esc(full_label(p['full'])) if full_label(p['full']) else ''
        out.append('<section class="slide" data-i="%d">%s<div class="lhs"><h2>%s</h2><div class="body">%s</div>%s</div>'
                   '<div class="rhs">%s</div></section>' % (i, kick, b3.inline(p['title']), render(p['body_md']), full,
                                                           render(p['fig_md'])))
    fonts = ('<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=DM+Sans:opsz,wght@9..40,400;9..40,500;9..40,700'
             '&family=DM+Mono:wght@400;500&family=Noto+Sans+SC:wght@400;500;700&display=swap">')
    return ('<!doctype html><html lang="%s"><head><meta charset="utf-8"><title>GEO Playbook %s</title>%s<style>%s\n%s</style></head><body>%s'
            '<script src="https://cdnjs.cloudflare.com/ajax/libs/mermaid/11.15.0/mermaid.min.js"></script><script>%s</script></body></html>') % (
        'zh-CN' if LANG == 'zh' else 'en', LANG, fonts, BOOK_CSS, SLIDE_CSS, ''.join(out), MMD_JS)

async def export(html_path, pdf_path, shots_dir):
    from playwright.async_api import async_playwright
    async with async_playwright() as pw:
        br = await pw.chromium.launch()
        pg = await br.new_page(viewport={'width': 1920, 'height': 1080}, device_scale_factor=2)
        await pg.goto('file://' + html_path, wait_until='networkidle')
        await pg.wait_for_function('window.__ready === true', timeout=60000)
        await pg.evaluate('document.fonts.ready')
        errs = await pg.evaluate("[...document.querySelectorAll('figure.mmd')].filter(f=>f.dataset.err).map(f=>f.dataset.err)")
        os.makedirs(shots_dir, exist_ok=True)
        slides = await pg.query_selector_all('section.slide')
        info = []
        for s in slides:
            i = await s.get_attribute('data-i')
            await s.screenshot(path=os.path.join(shots_dir, 'slide-%s.png' % i))
            fig = await s.query_selector('.rhs figure')
            fpath = ''
            if fig:
                fpath = os.path.join(shots_dir, 'fig-%s.png' % i)
                await fig.screenshot(path=fpath, omit_background=False)
            info.append({'i': int(i), 'fig': fpath})
        await pg.emulate_media(media='print')
        await pg.pdf(path=pdf_path, width='1920px', height='1080px', print_background=True, margin={'top': '0', 'bottom': '0', 'left': '0', 'right': '0'})
        await br.close()
        return info, errs

def pptx(pages, info, qr_bytes, out_path):
    from pptx import Presentation
    from pptx.util import Emu, Pt, Inches
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from PIL import Image
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    INK, INK2, MUTED, ACC, PAPER = RGBColor(0x1a, 0x19, 0x19), RGBColor(0x57, 0x53, 0x4e), RGBColor(0x6f, 0x6b, 0x66), RGBColor(0x4f, 0x46, 0xe5), RGBColor(0xfa, 0xfa, 0xf9)
    def bg(slide, color):
        f = slide.background.fill; f.solid(); f.fore_color.rgb = color
    def text(slide, x, y, w, h, s, size, color, bold=False, mono=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        for k, para in enumerate([q for q in s.split('\n') if q.strip()] or ['']):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = para; f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color
            f.name = 'Menlo' if mono else 'Arial'
            rpr = r._r.get_or_add_rPr(); ea = rpr.find(qn('a:ea'))
            if ea is None:
                ea = rpr.makeelement(qn('a:ea'), {}); rpr.append(ea)
            ea.set('typeface', 'PingFang SC')
            p.space_after = Pt(size * 0.5)
        return tb
    def plain(md):
        s = re.sub(r'```.*?```', '', md, flags=re.S)
        s = re.sub(r'\*\*([^*]+)\*\*', r'\1', s); s = re.sub(r'`([^`]+)`', r'\1', s)
        s = re.sub(r'\[\[([^\]]+)\]\]', r'\1', s); s = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', s)
        s = s.replace('[QR]', '')
        return '\n'.join(re.sub(r'^\s*[-*]\s+', '• ', l).strip() for l in s.split('\n') if l.strip() and not l.startswith('>'))
    figs = {x['i']: x['fig'] for x in info}
    n = len(pages)
    for i, p in enumerate(pages, 1):
        s = prs.slides.add_slide(blank)
        if i == 1:
            bg(s, INK)
            text(s, 0.9, 0.35, 11.5, 0.4, T['kick'], 11, RGBColor(0xa8, 0xa2, 0x9e), mono=True)
            text(s, 1.2, 2.0, 11, 2.6, p['title'], 44, PAPER, bold=True)
            text(s, 1.2, 4.6, 11, 1.6, plain(p['body_md']), 20, RGBColor(0xd6, 0xd3, 0xcd))
            text(s, 1.2, 6.4, 6, 0.5, 'canlah.ai', 14, RGBColor(0xb5, 0xbf, 0xfa), mono=True)
            continue
        bg(s, PAPER)
        text(s, 0.75, 0.3, 9, 0.4, T['kick'], 10, ACC, mono=True)
        text(s, 11.6, 0.3, 1.2, 0.4, '%d / %d' % (i, n), 10, MUTED, mono=True)
        # 标题按实际行数占位，正文紧跟其后（中文约 12 字一行，英文约 26 个字符一行）
        per = 12 if LANG != 'en' else 26
        lines = max(1, -(-len(p['title']) // per))
        text(s, 0.75, 1.2, 5.0, 0.6 * lines + 0.3, p['title'], 30, INK, bold=True)
        text(s, 0.75, 1.2 + 0.6 * lines + 0.45, 5.0, 3.6, plain(p['body_md']), 15, INK2)
        if p['qr']:
            qp = os.path.join(OUT, 'guide-%s-qr.png' % LANG); open(qp, 'wb').write(qr_bytes)
            s.shapes.add_picture(qp, Inches(8.6), Inches(1.5), Inches(3.4), Inches(3.4))
            text(s, 7.9, 5.0, 4.8, 0.5, URL.replace('https://', ''), 13, ACC, mono=True)
            text(s, 7.9, 5.45, 4.8, 0.5, T['scan'], 13, MUTED)
            text(s, 0.75, 6.55, 6.5, 0.6, T['agent'] + '：' + REPO, 11, INK2, mono=True)
            continue
        fp = figs.get(i)
        if fp and os.path.exists(fp):
            iw, ih = Image.open(fp).size
            maxw, maxh = 6.6, 5.6
            sc = min(maxw / iw, maxh / ih); w, h = iw * sc, ih * sc
            s.shapes.add_picture(fp, Inches(6.1 + (maxw - w) / 2), Inches(1.2 + (maxh - h) / 2), Inches(w), Inches(h))
        if p['full']:
            text(s, 0.75, 6.75, 5.2, 0.4, '→ ' + full_label(p['full']), 10, MUTED, mono=True)
    prs.save(out_path)

def web_json(pages):
    return [{'title': p['title'], 'fig': render(p['fig_md']), 'body': render(p['body_md']), 'full': p['full'], 'qr': p['qr']} for p in pages]

def main():
    src = os.path.join(V3, 'guide', '%s.md' % LANG)
    pages = parse(open(src).read())
    b3.WARN.clear()
    b3.LABEL_LANG = 'en' if LANG == 'en' else 'zh'
    qb = qr_png(URL); qb64 = base64.b64encode(qb).decode()
    os.makedirs(OUT, exist_ok=True)
    html_path = os.path.join(OUT, 'guide-%s.html' % LANG)
    open(html_path, 'w').write(slide_html(pages, qb64))
    pdf_path = os.path.join(OUT, 'GEO-Playbook-%s-2026.pdf' % ('轻量版' if LANG == 'zh' else 'Quick-Guide'))
    info, errs = asyncio.run(export(html_path, pdf_path, os.path.join(OUT, 'guide-%s-shots' % LANG)))
    pptx_path = pdf_path.replace('.pdf', '.pptx')
    pptx(pages, info, qb, pptx_path)
    json.dump(web_json(pages), open(os.path.join(OUT, 'guide-%s.json' % LANG), 'w'), ensure_ascii=False)
    print('%d 页 → %s\n       %s\n       %s' % (len(pages), html_path, pdf_path, pptx_path))
    print('二维码指向：', URL)
    if errs:
        print('⚠️ mermaid 渲染失败：', errs)
    if b3.WARN:
        print('⚠️', '\n'.join(b3.WARN))

main()
